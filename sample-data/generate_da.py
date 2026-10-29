# -*- coding: utf-8 -*-
"""Generator für Durchführungsaufträge (DA) je Projekt und Programmdurchführungs-
aufträge (PDA) je Programm der Stadt Musterstadt – HERMES-Struktur (Vorlage
durchfuehrungsauftrag.dotx).

Ablage spiegelt die PSR-Struktur (da-drive/), damit der rekursive SharePoint-Scan
später PSR UND DA findet und inhaltsbasiert unterscheidet. Inhalte sind konsistent
zu Portfolio-Soll (Budget/Aufwand/Status/Daten) und PSR-Stammdaten (AG/PL/Risiken).

Aufruf (Repo-Root):  python sample-data/generate_da.py
"""
import os
import shutil
import sys

import yaml
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from generate_psr import bold, chf, shade, tbl  # noqa: E402
from psr_data_demo import PROJECTS  # noqa: E402

OUT = os.path.join(HERE, "da-drive")

# Phasen-Aufteilung (Anteile) je Methode – für Kosten und Aufwand.
SPLIT = {
    "klassisch": [("Initialisierung", 0.10), ("Konzept (klassisch)", 0.20),
                  ("Realisierung (klassisch)", 0.40), ("Einführung (klassisch)", 0.20),
                  ("Abschluss", 0.10)],
    "agil": [("Initialisierung", 0.15), ("Umsetzung (agil)", 0.70), ("Abschluss", 0.15)],
}
# Vollständige Zeilenreihenfolge der Vorlage (Kosten/Aufwand je Phase).
PHASE_ROWS = ["Initialisierung", "Konzept (klassisch)", "Realisierung (klassisch)",
              "Einführung (klassisch)", "Umsetzung (agil)", "Abschluss"]
PROG_SPLIT = [("Programminitialisierung", 0.10), ("Programmdurchführung", 0.80),
              ("Programmabschluss", 0.10)]
STATUS_DE = {"idee": "Idee", "beantragt": "Beantragt", "freigegeben": "Freigegeben",
             "laufend": "In Durchführung", "zurueckgestellt": "Zurückgestellt",
             "abgeschlossen": "Abgeschlossen", "abgebrochen": "Abgebrochen"}


def _safe(name):
    for ch in '<>:"/\\|?*':
        name = name.replace(ch, "")
    return name.strip()


def kuerzel(name):
    parts = [p for p in name.replace(".", " ").split() if p]
    if not parts:
        return ""
    return (parts[0][0] + parts[-1][:2]).upper() if len(parts) > 1 else parts[0][:3].upper()


def de(d):  # '2024-01-15' -> '15.01.2024'
    y, m, day = (d.split("-") + ["01", "01"])[:3]
    return f"{day}.{m}.{y}"


def split_values(total, method):
    out = {ph: "" for ph in PHASE_ROWS}
    acc = 0
    defs = SPLIT.get(method, SPLIT["klassisch"])
    for ph, prop in defs:
        v = round(total * prop)
        out[ph] = v
        acc += v
    return out, acc


def meta_table(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text = k
        shade(cells[0], "F2F2F2")
        bold(cells[0])
        cells[1].text = str(v)


def kosten_tabelle(doc, titel, einheit, werte, total, suffix=""):
    doc.add_heading(titel, level=3)
    rows = [(ph, (f"{chf(werte[ph])}{suffix}" if werte[ph] != "" else "–")) for ph in PHASE_ROWS]
    rows.append(("Total", f"{chf(total)}{suffix}"))
    tbl(doc, ["Phase", einheit], rows)


def risiko_rows(risks, pl):
    out = []
    sev = [(2, 2), (3, 2), (2, 3)]
    for i, r in enumerate(risks):
        ew, ag = sev[i % len(sev)]
        out.append((f"R{i+1}", r, ew, ag, ew * ag,
                    "Gegensteuernde Massnahme definiert", pl, ""))
    return out


# ---- Projekt-DA -----------------------------------------------------------
def build_projekt_da(p, soll, zeit, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    a = soll.get("auftrag", {})
    z = zeit or {}
    method = z.get("method", "klassisch")
    status = a.get("status", "freigegeben")
    budget = a.get("budget_soll", p["cost_plan"])
    aufwand = a.get("aufwand_soll", p["eff_plan"])

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Durchführungsauftrag", level=0)

    meta_table(doc, [
        ("Projektname", p["name"]), ("Projektnummer", p["nummer"]),
        ("Auftraggeber/in", p["ag"]), ("Projektleiter/in", p["pl"]),
        ("Verwaltungseinheit", p["ve"]), ("Geschäftsbereich", p["gb"]),
        ("Programm", soll.get("programm", "–") or "–"),
        ("Vorgehensmodell", "Agil (Scrum)" if method == "agil" else "Klassisch (HERMES)"),
        ("Geplanter Start", de(z.get("start", a.get("datum", "")))),
        ("Geplantes Ende", de(z.get("ende", ""))),
        ("Status", STATUS_DE.get(status, status)),
    ])

    doc.add_heading("Ausgangslage und Handlungsbedarf", level=1)
    doc.add_paragraph(f"Das Vorhaben «{p['name']}» adressiert {p['ziel'][1].lower()}. "
                      "Die heutige Lösung genügt den Anforderungen nicht mehr; es besteht Handlungsbedarf.")

    doc.add_heading("Ziele", level=1)
    doc.add_paragraph(f"Lösungsziel ({p['ziel'][0]}): {p['ziel'][1]}.")

    doc.add_heading("Investitionsbedarf", level=1)
    kw, kt = split_values(budget, method)
    kosten_tabelle(doc, "Projektkosten", "Geplant CHF", kw, kt)
    aw, at = split_values(aufwand, method)
    kosten_tabelle(doc, "Arbeitstage", "Geplant Arbeitstage", aw, at)

    doc.add_heading("Planung und Organisation", level=1)
    doc.add_heading("Projektorganisation", level=2)
    tbl(doc, ["Rolle in der Projektorganisation", "Name", "Kürzel", "Funktion / Vertretene Organisation"],
        [("Auftraggeber/in", p["ag"], kuerzel(p["ag"]), p["gb"]),
         ("Projektleiter/in", p["pl"], kuerzel(p["pl"]), p["ve"]),
         ("Fachverantwortung", "N. N.", "", p["ve"]),
         ("Qualitätssicherung", "N. N.", "", "Stabsstelle PMO")])

    doc.add_heading("Risiken", level=1)
    tbl(doc, ["Nr.", "Risikobeschreibung", "EW", "AG", "RZ", "Massnahmen", "Verantwortlich", "Termin"],
        risiko_rows(p["risks"], p["pl"]))

    doc.add_heading("Konsequenzen", level=1)
    tbl(doc, ["Zeitpunkt", "Beschreibung"],
        [("Bei Durchführungsfreigabe", f"{p['ziel'][1]} kann termingerecht umgesetzt werden."),
         ("Bei verspäteter Durchführungsfreigabe", "Verzögerung mit Mehrkosten; der Nutzen tritt später ein."),
         ("Bei ausbleibender / abgelehnter Freigabe", "Handlungsbedarf bleibt ungelöst; Ziel wird nicht erreicht.")])

    doc.add_heading("Verbindlichkeiten", level=1)
    tbl(doc, ["Nr.", "Verbindlichkeit", "Beschreibung", "Verantwortlich"],
        [("01", "Budget", f"Das genehmigte Budget von {chf(budget)} CHF ist einzuhalten.", p["pl"]),
         ("02", "Berichterstattung", "Monatlicher Projektstatusbericht (PSR) an den Auftraggeber.", p["pl"])])

    doc.add_heading("Dokument-Protokoll", level=1)
    freigegeben = status in ("freigegeben", "laufend", "abgeschlossen")
    tbl(doc, ["Version", "Name", "Datum", "Bemerkungen"],
        [("1.0", p["ag"], de(a.get("datum", "")),
          "Durchführung freigegeben" if freigegeben else "Entwurf zur Prüfung")])

    doc.save(path)


# ---- Programm-DA (PDA) ----------------------------------------------------
def build_programm_da(name, ref, soll, zeit, kinder, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    a = soll.get("auftrag", {})
    z = zeit or {}
    budget = a.get("budget_soll", 0)
    aufwand = a.get("aufwand_soll", 0)
    status = a.get("status", "laufend")

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Programmdurchführungsauftrag", level=0)

    meta_table(doc, [
        ("Programmname", name), ("Programmnummer", ref),
        ("Auftraggeber/in", "Stadtrat Musterstadt"),
        ("Programmleitung", "N. N."),
        ("Geplanter Start", de(z.get("start", a.get("datum", "")))),
        ("Geplantes Ende", de(z.get("ende", ""))),
        ("Status", STATUS_DE.get(status, status)),
    ])

    doc.add_heading("Ausgangslage und Handlungsbedarf", level=1)
    doc.add_paragraph(f"Das Programm «{name}» bündelt {len(kinder)} Projekte zu einem "
                      "koordinierten Veränderungsvorhaben der Stadtverwaltung.")

    doc.add_heading("Investitionsbedarf", level=1)
    doc.add_heading("Programmkosten", level=3)
    rows, acc = [], 0
    for ph, prop in PROG_SPLIT:
        v = round(budget * prop)
        acc += v
        rows.append((ph, f"{chf(v)}"))
    rows.append(("Total", chf(acc)))
    tbl(doc, ["Phase", "Geplant CHF"], rows)
    doc.add_heading("Arbeitstage", level=3)
    rows = []
    for ph, prop in PROG_SPLIT:
        rows.append((ph, chf(round(aufwand * prop))))
    rows.append(("Total", chf(aufwand)))
    tbl(doc, ["Phase", "Geplant Arbeitstage"], rows)

    doc.add_heading("Enthaltene Projekte", level=1)
    doc.add_paragraph("Die folgenden Projekte werden innerhalb der Programmdurchführung geführt:")
    tbl(doc, ["Nr.", "Projekt", "Budget CHF", "Status"],
        [(k["nummer"], k["name"], chf(k["budget"]), STATUS_DE.get(k["status"], k["status"]))
         for k in kinder] or [("–", "", "", "")])

    doc.add_heading("Programmorganisation", level=1)
    tbl(doc, ["Rolle in der Programmorganisation", "Name", "Kürzel", "Funktion / Vertretene Organisation"],
        [("Auftraggeber/in", "Stadtrat Musterstadt", "STR", "Stadtrat"),
         ("Programmleitung", "N. N.", "", "Stabsstelle PMO"),
         ("Programmausschuss", "N. N.", "", "Ressortleitungen")])

    doc.add_heading("Konsequenzen", level=1)
    tbl(doc, ["Zeitpunkt", "Beschreibung"],
        [("Bei Durchführungsfreigabe", "Die Projekte des Programms können koordiniert umgesetzt werden."),
         ("Bei ausbleibender Freigabe", "Die Digitalisierungsziele des Programms werden nicht erreicht.")])

    doc.add_heading("Dokument-Protokoll", level=1)
    tbl(doc, ["Version", "Name", "Datum", "Bemerkungen"],
        [("1.0", "Stadtrat Musterstadt", de(a.get("datum", "")), "Programmdurchführung freigegeben")])

    doc.save(path)


def _stub(ref, v):
    """Minimale Stammdaten für Portfolio-Projekte ohne PSR-Stammdaten (z. B. geplant)."""
    a = v.get("auftrag", {})
    return {"short": ref, "name": v.get("titel") or ref, "nummer": ref,
            "pl": "N. N.", "ag": "Ressortleitung " + (v.get("bereich", "").replace("Ressort ", "") or "–"),
            "ve": v.get("einheit", "–"), "gb": v.get("bereich", "–"),
            "cost_plan": a.get("budget_soll", 0), "eff_plan": a.get("aufwand_soll", 0),
            "ziel": ("Nutzen", v.get("titel", "das Vorhaben") + " umsetzen"),
            "risks": ["Datenübernahme aus Altsystem", "Verfügbarkeit Fachressourcen"]}


def main():
    pf = yaml.safe_load(open(os.path.join(HERE, "portfolio-musterstadt.yaml"), encoding="utf-8"))
    soll = {v["ref"]: v for v in pf["vorhaben"]}
    zeit = pf.get("zeitplan", {})
    by_short = {p["nummer"]: p for p in PROJECTS}

    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    os.makedirs(OUT)

    # Programm-Kinder sammeln (für die PDA-Projektliste)
    kinder = {}
    for ref, v in soll.items():
        if v.get("art") == "projekt" and v.get("programm"):
            p = by_short.get(ref)
            kinder.setdefault(v["programm"], []).append({
                "nummer": ref, "name": v.get("titel") or ref,
                "budget": v["auftrag"].get("budget_soll", 0),
                "status": v["auftrag"].get("status", "")})

    n_proj = n_prog = 0
    for ref, v in soll.items():
        if v.get("art") == "programm":
            name = v.get("titel") or ref
            path = os.path.join(OUT, "Programme", _safe(name), "Programmdurchfuehrungsauftrag.docx")
            build_programm_da(name, ref, v, zeit.get(ref), kinder.get(name, []), path)
            n_prog += 1
        else:
            p = by_short.get(ref) or _stub(ref, v)
            prog = v.get("programm")
            if prog:
                folder = os.path.join(OUT, "Programme", _safe(prog), _safe(p["name"]))
            else:
                folder = os.path.join(OUT, "Fachbereiche", _safe(p["gb"]), _safe(p["name"]))
            build_projekt_da(p, v, zeit.get(ref), os.path.join(folder, "Durchfuehrungsauftrag.docx"))
            n_proj += 1

    print(f"Durchführungsaufträge erzeugt in {OUT}")
    print(f"  {n_proj} Projekt-DA, {n_prog} Programm-DA")


if __name__ == "__main__":
    main()
