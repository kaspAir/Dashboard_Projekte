# -*- coding: utf-8 -*-
"""Zeitreihen-Generator: pro Projekt monatliche PSR (docx) + ground_truth.yaml.

Ordnerstruktur:  sample-data/psr/<NN>_<Short>/<YYYY-MM>_PSR.docx
Oracle:          sample-data/ground_truth.yaml  (alle Snapshots als Zeitreihe)
"""
import os
import calendar
import shutil
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from psr_data import PROJECTS

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "psr")
CURRENT = (2026, 3)        # "heute": letzter generierter Berichtsmonat
WINDOW_START = (2024, 1)   # frueheste generierte Periode (kappt lange Historien)

SK = ["gesamt", "termine", "kosten", "personalaufwand",
      "ergebnisse", "projektziele", "projektrisiken", "governance"]
BEREICHE = [("Gesamtbeurteilung", "gesamt"), ("Termine", "termine"), ("Kosten", "kosten"),
            ("Personalaufwand", "personalaufwand"), ("Ergebnisse", "ergebnisse"),
            ("Projektziele", "projektziele"), ("Projektrisiken", "projektrisiken"),
            ("Einhaltung Governance", "governance")]
PHASES = [(0.12, "Initialisierung"), (0.32, "Konzept"), (0.75, "Realisierung"),
          (0.95, "Einführung"), (1.01, "Abschluss")]
MILESTONES = [(0.30, "Konzept abgeschlossen"), (0.60, "Realisierung Kernfunktion"),
              (0.90, "Einführung / Go-Live")]
OVERRUN = {"gescheitert_frueh": 1.15, "gescheitert_spaet": 1.22,
           "verschlechterung": 1.20, "volatil": 1.06, "gelb_stabil": 1.05}
# Vorwarnzeit bei Abbruch: (Monate rot, zusätzliche Monate gelb davor) vor Projektende.
# Kurze Werte = "spät gewarnt" (unehrliche Berichterstattung).
ABORT_LEAD = {"EPatientenakte": (6, 3), "Auslaenderregister": (4, 3), "HRSystem": (1, 1)}
RAG_FILL = {"grün": "C6EFCE", "gelb": "FFEB9C", "rot": "FFC7CE"}

# ---- Zeit-Utilities --------------------------------------------------------
def midx(ym):
    return ym[0] * 12 + (ym[1] - 1)

def add_months(ym, d):
    x = midx(ym) + d
    return (x // 12, x % 12 + 1)

def month_end(ym):
    y, m = ym
    return f"{calendar.monthrange(y, m)[1]:02d}.{m:02d}.{y}"

def tag(ym):
    return f"{ym[0]:04d}-{ym[1]:02d}"

def chf(n):
    return format(int(round(n)), ",").replace(",", "'")

# ---- Ampel-/Kennzahl-Logik -------------------------------------------------
def phase_for(p):
    for thr, name in PHASES:
        if p < thr:
            return name
    return "Abschluss"

def overall(arch, p, i):
    if arch == "stabil":
        return "grün"
    if arch == "solide":
        return "gelb" if 0.42 <= p < 0.58 else "grün"
    if arch == "gelb_stabil":
        return "grün" if p < 0.15 else "gelb"
    if arch == "erholung":
        return "rot" if p < 0.30 else ("gelb" if p < 0.60 else "grün")
    if arch == "verschlechterung":
        return "grün" if p < 0.40 else ("gelb" if p < 0.70 else "rot")
    if arch == "volatil":
        return ["grün", "gelb", "grün", "rot", "gelb", "grün", "gelb", "rot", "gelb"][(i // 3) % 9]
    if arch == "gescheitert_frueh":
        return "grün" if p < 0.25 else ("gelb" if p < 0.45 else "rot")
    if arch == "gescheitert_spaet":
        return "grün" if p < 0.60 else ("gelb" if p < 0.82 else "rot")
    return "grün"

def statuses(ov, weak):
    st = {k: "grün" for k in SK}
    st["gesamt"] = ov
    if ov == "gelb":
        for w in (weak[:2] or ["termine"]):
            st[w] = "gelb"
    elif ov == "rot":
        for w in (weak or ["termine"]):
            st[w] = "rot"
        if st["projektrisiken"] == "grün":
            st["projektrisiken"] = "gelb"
    return st

def erlaeuterung(st, weak, ov):
    hints = {"termine": "Terminverzug gegenüber Plan", "kosten": "Kosten über Plan",
             "personalaufwand": "Ressourcenengpass im Kernteam", "ergebnisse": "Ergebnis nicht abnahmefähig",
             "projektziele": "Zielerreichung gefährdet", "projektrisiken": "hohe Risiken offen",
             "governance": "Vorgaben nur teilweise erfüllt"}
    return {k: hints[k] for k in SK if st[k] != "grün" and k != "gesamt"}

def fazit(ov, lifecycle, is_last):
    if is_last and lifecycle == "completed":
        return "Projekt erfolgreich abgeschlossen; die Ziele wurden im Wesentlichen erreicht."
    if is_last and lifecycle == "aborted":
        return "Projekt abgebrochen: Erfolg nicht mehr erreichbar, Abbruch durch den Auftraggeber entschieden."
    return {"grün": "Projekt auf Kurs; im Zeit- und Kostenrahmen.",
            "gelb": "Projekterfolg gefährdet, aber durch die Projektorganisation steuerbar.",
            "rot": "Projekterfolg stark gefährdet; Eskalation an die Stammorganisation."}[ov]

# ---- docx-Bau --------------------------------------------------------------
def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexc)
    tcPr.append(shd)

def bold(cell):
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True

def tbl(doc, headers, rows, rag_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        shade(c, "D9E2F3")
        bold(c)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
            if rag_col is not None and i == rag_col and str(v) in RAG_FILL:
                shade(cells[i], RAG_FILL[str(v)])

def build_doc(s, path):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Projektstatusbericht", level=0)

    m = doc.add_table(rows=0, cols=2)
    m.style = "Table Grid"
    for k, v in [("Projektname", s["name"]), ("Projektnummer", s["nummer"]),
                 ("Bearbeitungsdatum", s["datum"]), ("Berichtsperiode", s["periode"]),
                 ("Projektleiter/in", s["pl"]), ("Auftraggeber/in", s["ag"]),
                 ("Verwaltungseinheit", s["ve"]), ("Geschäftsbereich", s["gb"]),
                 ("Projektphase", s["phase"]), ("Projektstatus (Lebenszyklus)", s["lifecycle_de"])]:
        cells = m.add_row().cells
        cells[0].text = k
        shade(cells[0], "F2F2F2")
        bold(cells[0])
        cells[1].text = v

    doc.add_heading("Übersicht Projektstand", level=1)
    doc.add_heading("Gesamtstatus im Überblick", level=2)
    tbl(doc, ["Bereich", "Status", "Erläuterungen wenn nicht grün"],
        [(lab, s["status"][key], s["erl"].get(key, "")) for lab, key in BEREICHE], rag_col=1)

    doc.add_heading("Gesamtbeurteilung / Fazit", level=2)
    doc.add_paragraph(s["fazit"])

    doc.add_heading("Prognose der Zielerreichung", level=2)
    tbl(doc, ["Nr.", "Kategorie", "Zielformulierung", "Prognose Zielerreichung"],
        [("LZ-01", s["ziel"][0], s["ziel"][1], s["ziel_pct"])])

    doc.add_heading("Soll/Ist-Vergleich und Prognosen", level=2)
    doc.add_heading("Kosten (CHF)", level=3)
    tbl(doc, ["Position", "Ist", "Plan gesamt", "Prognose gesamt", "Differenz"],
        [("Total", s["k_ist"], s["k_plan"], s["k_prog"], s["k_diff"])])
    doc.add_heading("Aufwand (Personentage)", level=3)
    tbl(doc, ["Position", "Ist", "Plan gesamt", "Prognose gesamt", "Differenz"],
        [("Total", s["a_ist"], s["a_plan"], s["a_prog"], s["a_diff"])])

    doc.add_heading("Erreichte Meilensteine (vergangene Berichtsperioden)", level=2)
    tbl(doc, ["Meilenstein", "Plan", "Prognose", "Ist"], s["ms_erreicht"] or [("–", "", "", "")])
    doc.add_heading("Geplante Meilensteine (nächste Berichtsperioden)", level=2)
    tbl(doc, ["Meilenstein", "Plan", "Prognose", "Ist"], s["ms_geplant"] or [("–", "", "", "")])

    if s["probleme"]:
        doc.add_heading("Probleme und Massnahmen", level=2)
        tbl(doc, ["Nr.", "Problembeschreibung", "Massnahme", "Verantwortlich", "Termin"], s["probleme"])

    doc.add_heading("Risiken", level=2)
    tbl(doc, ["Nr.", "Risikobeschreibung", "EW", "AG", "RZ", "Verantwortlich"], s["risiken"])

    doc.save(path)

# ---- Simulation ------------------------------------------------------------
LIFEC_DE = {"active": "aktiv (laufend)", "completed": "abgeschlossen", "aborted": "abgebrochen"}

def snapshots_for(p):
    start = p["start"]
    dur = p["dur"]
    end = CURRENT if p["lifecycle"] == "active" else add_months(start, dur)
    if midx(end) > midx(CURRENT):
        end = CURRENT
    series_start = start if midx(start) >= midx(WINDOW_START) else WINDOW_START
    months = []
    cur = series_start
    while midx(cur) <= midx(end):
        months.append(cur)
        cur = add_months(cur, 1)
    over = OVERRUN.get(p["arch"], 1.0)
    k_forecast_total = p["cost_plan"] * over
    a_forecast_total = p["eff_plan"] * over
    ziel_target = 100 if p["arch"] in ("stabil", "solide", "erholung", "gelb_stabil") else 60

    out = []
    for i, ym in enumerate(months):
        is_last = (ym == end)
        prog = (midx(ym) - midx(start) + 1) / dur
        prog = max(0.03, min(prog, 1.0))
        ov = overall(p["arch"], prog, i)
        if p["lifecycle"] == "aborted":
            red_lead, amber_extra = ABORT_LEAD.get(p["short"], (5, 3))
            mte = midx(end) - midx(ym)
            ov = "rot" if mte < red_lead else ("gelb" if mte < red_lead + amber_extra else "grün")
        if is_last and p["lifecycle"] == "completed":
            ov = "grün"
        st = statuses(ov, p["weak"])
        phase = "Abschluss" if (is_last and p["lifecycle"] == "completed") else phase_for(prog)

        k_ist = k_forecast_total * prog
        a_ist = a_forecast_total * prog
        # Meilensteine
        ms_e, ms_g = [], []
        for frac, nm in MILESTONES:
            ms_date = month_end(add_months(start, int(round(frac * dur))))
            if prog >= frac:
                ms_e.append((nm, ms_date, ms_date, ms_date))
            else:
                ms_g.append((nm, ms_date, ms_date, ""))
        # Risiken (mehr/heftiger bei gelb/rot)
        sev = {"grün": (1, 2), "gelb": (2, 2), "rot": (3, 3)}[ov]
        nrisk = 1 if ov == "grün" else 2
        risiken = [(f"R{j+1}", p["risks"][j], sev[0], sev[1], sev[0] * sev[1], "PL")
                   for j in range(min(nrisk, len(p["risks"])))]
        # Probleme nur bei gelb/rot
        probleme = []
        if ov in ("gelb", "rot"):
            wk = (p["weak"] or ["termine"])[0]
            probleme = [("P1", erlaeuterung(st, p["weak"], ov).get(wk, "Abweichung gegenüber Plan"),
                         "Gegensteuernde Massnahme eingeleitet", "PL", month_end(add_months(ym, 1)))]

        rec = dict(
            name=p["name"], nummer=p["nummer"], datum=month_end(ym), periode=tag(ym),
            pl=p["pl"], ag=p["ag"], ve=p["ve"], gb=p["gb"], phase=phase,
            lifecycle=p["lifecycle"], lifecycle_de=(LIFEC_DE[p["lifecycle"]] if is_last or p["lifecycle"] == "active" else "aktiv (laufend)"),
            status=st, erl=erlaeuterung(st, p["weak"], ov),
            fazit=fazit(ov, p["lifecycle"], is_last),
            ziel=p["ziel"], ziel_pct=f"{int(round(min(prog,1.0)*ziel_target))}%",
            k_plan=chf(p["cost_plan"]), k_ist=chf(k_ist), k_prog=chf(k_forecast_total),
            k_diff=("+" if k_forecast_total > p["cost_plan"] else "") + chf(k_forecast_total - p["cost_plan"]),
            a_plan=chf(p["eff_plan"]), a_ist=chf(a_ist), a_prog=chf(a_forecast_total),
            a_diff=("+" if a_forecast_total > p["eff_plan"] else "") + chf(a_forecast_total - p["eff_plan"]),
            ms_erreicht=ms_e, ms_geplant=ms_g, probleme=probleme, risiken=risiken,
        )
        # nur der letzte Snapshot eines inaktiven Projekts traegt den Endzustand
        if p["lifecycle"] != "active":
            rec["lifecycle_de"] = LIFEC_DE[p["lifecycle"]] if is_last else "aktiv (laufend)"
        out.append(rec)
    return out


if __name__ == "__main__":
    import yaml
    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    os.makedirs(OUT)
    manifest = []
    total = 0
    for idx, p in enumerate(PROJECTS, start=1):
        folder = os.path.join(OUT, f"{idx:02d}_{p['short']}")
        os.makedirs(folder, exist_ok=True)
        for s in snapshots_for(p):
            fname = f"{s['periode']}_PSR.docx"
            build_doc(s, os.path.join(folder, fname))
            total += 1
            manifest.append(dict(
                projekt=p["short"], name=p["name"], nummer=p["nummer"],
                datei=f"{idx:02d}_{p['short']}/{fname}", periode=s["periode"],
                verwaltungseinheit=p["ve"], geschaeftsbereich=p["gb"],
                lifecycle=p["lifecycle"], phase=s["phase"],
                status=s["status"],
                kosten={"plan": s["k_plan"], "ist": s["k_ist"], "prognose": s["k_prog"]}))
        print(f"{idx:02d} {p['short']:18s} {p['lifecycle']:9s} {len(snapshots_for(p)):2d} Snapshots")

    with open(os.path.join(HERE, "ground_truth.yaml"), "w", encoding="utf-8") as f:
        yaml.safe_dump(manifest, f, allow_unicode=True, sort_keys=False)
    print(f"\nTotal {total} PSR-Snapshots in {len(PROJECTS)} Projektordnern + ground_truth.yaml")
