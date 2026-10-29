# -*- coding: utf-8 -*-
"""Erzeugt ein realistisches „Projektlaufwerk" für den Demo-Mandanten Stadt Musterstadt.

Zweck der Demo: Das Dashboard bekommt NUR den Wurzelordner (bzw. die SharePoint-Site)
und findet die Statusberichte selbst – egal wie tief verschachtelt –, erkennt sie
inhaltsbasiert und ignoriert alles, was kein PSR ist.

Struktur:
  musterstadt-drive/
    Programme/<Programm>/<Projekt>/Statusberichte/<YYYY-MM>_PSR.docx
    Fachbereiche/<Ressort>/<Projekt>/Berichte/<YYYY-MM>_PSR.docx
    + verstreute Nicht-PSR-Dateien (Protokolle, Konzepte, Angebote, Richtlinien)

Projekte sind durchgehend klassisch ODER agil (nie gemischt). Ein paar Projekte
sind zu Programmen gebündelt, andere stehen eigenständig.

Aufruf (Repo-Root):  python sample-data/generate_demo_drive.py
"""
import os
import shutil
import sys

import yaml
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.dirname(HERE))

from generate_psr import add_months, build_doc, midx, snapshots_for, tag  # noqa: E402
from psr_data_demo import PROJECTS  # noqa: E402
from app.ingestion import run_ingestion  # noqa: E402

DRIVE = os.path.join(HERE, "musterstadt-drive")

# --- Vorgehensmodell je Projekt (durchgehend, nie gemischt) ---------------
AGILE = {"Buergerportal", "SozialPortal", "Einsatzapp", "Ratsinfo", "BiblioCloud"}

# --- Programme (Programm -> seine Projekte); der Rest ist eigenständig -----
PROGRAMME = {
    "Digitale Verwaltung Musterstadt": ["Buergerportal", "Baugenehmigung", "ERechnung", "Ratsinfo"],
    "Smart City Infrastruktur": ["Verkehrsleit", "Umweltsensorik", "SchulWLAN"],
    "Gesundheit & Soziales digital": ["SozialPortal", "GesundheitTermin"],
}
PROJ_PROGRAMM = {short: pg for pg, shorts in PROGRAMME.items() for short in shorts}

AGILE_PHASE = {"Initialisierung": "Vorbereitung", "Konzept": "Vorbereitung",
               "Realisierung": "Umsetzung (Sprints)", "Einführung": "Einführung",
               "Abschluss": "Abschluss"}

SPRINT_THEMES = ["Setup & Architektur", "Kernfunktion", "Datenmodell", "Schnittstellen",
                 "Self-Service", "Benachrichtigungen", "Rollen & Rechte", "Reporting",
                 "Barrierefreiheit", "Stabilisierung", "Migration", "Abnahmetests"]


def _safe(name):
    for ch in '<>:"/\\|?*':
        name = name.replace(ch, "")
    return name.strip()


def sprints_for(p, periode, ov, is_last):
    """Sprint-Zeilen (abgeschlossen, geplant) für einen agilen Snapshot."""
    y, mth = (int(x) for x in periode.split("-"))
    me = midx((y, mth)) - midx(p["start"])
    cur = max(1, me * 2)                                  # ~2 Sprints/Monat, laufender Sprint
    base = max(10, int(round(p["eff_plan"] / (p["dur"] * 2) / 5)) * 5)   # Plan-Velocity
    factor = {"grün": 1.0, "gelb": 0.82, "rot": 0.6}[ov]

    def span(n):
        mon = add_months(p["start"], (n - 1) // 2)
        half = "1. Hälfte" if (n - 1) % 2 == 0 else "2. Hälfte"
        return f"{tag(mon)} {half}"

    def ist(n):
        return int(round(base * factor * (0.85 + 0.3 * ((n * 7) % 4) / 3)))

    done = [(f"Sprint {n}", span(n), SPRINT_THEMES[(n - 1) % len(SPRINT_THEMES)],
             f"Plan {base} / Ist {ist(n)}")
            for n in range(max(1, cur - 2), cur + 1)]
    nxt = []
    if not (is_last and p["lifecycle"] in ("completed", "aborted")):
        nxt = [(f"Sprint {n}", span(n), SPRINT_THEMES[(n - 1) % len(SPRINT_THEMES)], f"Plan {base}")
               for n in range(cur + 1, cur + 3)]
    return done, nxt


def write_psr(p, out_folder):
    os.makedirs(out_folder, exist_ok=True)
    method = "agil" if p["short"] in AGILE else "klassisch"
    programm = PROJ_PROGRAMM.get(p["short"], "")
    snaps = snapshots_for(p)
    for i, s in enumerate(snaps):
        s["method"] = method
        s["programm"] = programm
        if method == "agil":
            s["phase"] = AGILE_PHASE.get(s["phase"], s["phase"])
            done, nxt = sprints_for(p, s["periode"], s["status"]["gesamt"], i == len(snaps) - 1)
            s["sp_erreicht"], s["sp_geplant"] = done, nxt
        build_doc(s, os.path.join(out_folder, f"{s['periode']}_PSR.docx"))
    return len(snaps)


# --- Nicht-PSR-Dateien (Köder) --------------------------------------------
def decoy(path, title, paras, with_table=False):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d = Document()
    d.add_heading(title, level=1)
    for para in paras:
        d.add_paragraph(para)
    if with_table:                       # bewusst eine NICHT-Status-Tabelle (Falle)
        t = d.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        for i, h in enumerate(["Teilnehmer", "Funktion", "Anwesend"]):
            t.rows[0].cells[i].text = h
        for r in [("A. Meier", "Leitung", "ja"), ("B. Huber", "Protokoll", "ja")]:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = v
    d.save(path)


DECOYS = [
    ("Ablagerichtlinie_Projektlaufwerk.docx", "Ablagerichtlinie Projektlaufwerk",
     ["Diese Richtlinie regelt die Ablage von Projektdokumenten der Stadt Musterstadt.",
      "Statusberichte sind im jeweiligen Projektordner abzulegen."], False),
    ("Programme/Digitale Verwaltung Musterstadt/Programmhandbuch.docx", "Programmhandbuch",
     ["Das Programm bündelt die Digitalisierungsvorhaben der Stadtverwaltung.",
      "Steuerung durch den Programmausschuss."], False),
    ("Programme/Digitale Verwaltung Musterstadt/Buergerportal/Grobkonzept.docx",
     "Grobkonzept Bürgerportal", ["Fachliches Grobkonzept, Version 0.9 (Entwurf)."], True),
    ("Programme/Digitale Verwaltung Musterstadt/Buergerportal/Statusberichte/Lesemich.docx",
     "Hinweise zur Ablage", ["PSR bitte monatlich hochladen."], False),
    ("Programme/Smart City Infrastruktur/Verkehrsleit/Angebot_Lieferant.docx",
     "Angebot Sensorik-Lieferant", ["Angebot über Lieferung und Montage der Verkehrssensorik."], True),
    ("Fachbereiche/Ressort Sicherheit/Einwohnerregister-Ablösung/Sitzungsprotokoll_2024-11.docx",
     "Protokoll Lenkungsausschuss", ["Protokoll der Sitzung des Lenkungsausschusses.",
      "Beschluss: Projektabbruch wird geprüft."], True),
    ("Fachbereiche/Ressort Soziales & Gesundheit/Schulungsunterlagen.docx",
     "Schulungsunterlagen Fachanwendung", ["Foliensatz für die Anwenderschulung."], False),
    ("Fachbereiche/Ressort Bau & Umwelt/Umwelt-Sensornetz/Wartungsvertrag.docx",
     "Wartungsvertrag Sensorik", ["Vertrag über die Wartung des Sensornetzes."], True),
]


def main():
    if os.path.isdir(DRIVE):
        shutil.rmtree(DRIVE)
    os.makedirs(DRIVE)

    total_psr = 0
    for idx, p in enumerate(PROJECTS, start=1):
        pname = _safe(p["name"])
        programm = PROJ_PROGRAMM.get(p["short"])
        if programm:
            folder = os.path.join(DRIVE, "Programme", _safe(programm), pname, "Statusberichte")
        else:
            folder = os.path.join(DRIVE, "Fachbereiche", _safe(p["gb"]), pname, "Berichte")
        total_psr += write_psr(p, folder)

    for rel, title, paras, with_table in DECOYS:
        decoy(os.path.join(DRIVE, rel), title, paras, with_table)

    print(f"Laufwerk erzeugt: {DRIVE}")
    print(f"  {len(PROJECTS)} Projekte, {total_psr} PSR, {len(DECOYS)} Nicht-PSR-Dateien")

    # Store aus dem Laufwerk extrahieren (genau wie später vom SharePoint) ----
    result = run_ingestion({"type": "folder", "path": DRIVE})
    out = os.path.join(HERE, "store-musterstadt.yaml")
    with open(out, "w", encoding="utf-8") as f:
        yaml.safe_dump(result["records"], f, allow_unicode=True, sort_keys=False)
    print(f"\nRoot-Scan: gescannt {result['scanned']}, als PSR erkannt {result['ingested']}, "
          f"übersprungen {result['skipped']}")
    print(f"Store: {out} ({len(result['records'])} Snapshots)")


if __name__ == "__main__":
    main()
