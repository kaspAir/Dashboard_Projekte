"""Rohzugriff auf ein .docx: Tabellen als Listen von Zellen-Strings."""
from __future__ import annotations

from docx import Document


def read_tables(path: str):
    doc = Document(path)
    return [[[cell.text.strip() for cell in row.cells] for row in t.rows] for t in doc.tables]


def read_paragraphs(path: str):
    """Alle nichtleeren Absätze (Fliesstext) in Dokumentreihenfolge."""
    doc = Document(path)
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def text_after_heading(paragraphs, heading: str) -> str:
    """Erster Absatz nach einer Überschrift (Überschrift per Teilstring erkannt)."""
    h = heading.lower()
    seen = False
    for t in paragraphs:
        if seen:
            return t
        if h in t.lower():
            seen = True
    return ""


def metadata_kv(tables):
    """Die Metadaten-Tabelle (Key/Value) anhand der Zeile 'Projektname' finden."""
    for t in tables:
        for row in t:
            if row and row[0] == "Projektname":
                return {r[0]: (r[1] if len(r) > 1 else "") for r in t if r and r[0]}
    return {}


def tables_with_header(tables, keywords):
    """Alle Tabellen, deren Kopfzeile alle Stichwörter enthält (Reihenfolge erhalten)."""
    res = []
    for t in tables:
        if not t:
            continue
        header = " | ".join(t[0]).lower()
        if all(k.lower() in header for k in keywords):
            res.append(t)
    return res
